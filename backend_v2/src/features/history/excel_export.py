"""
Excel Export Deep Module — Generates professional styled Excel workbooks for carton history & traceability.
Supports Summary mode (single sheet) and Detailed Traceability mode (2-sheet workbook).
"""
import io

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.core import models


def style_header_row(ws, headers: list[str], header_fill_color: str = "1E1B4B"):
    """Styles the header row with dark background, bold white text, and borders."""
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color=header_fill_color, end_color=header_fill_color, fill_type="solid")
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC")
    )

    ws.append(headers)
    ws.row_dimensions[1].height = 28

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = alignment
        cell.border = thin_border


def auto_fit_columns(ws, max_cols: int):
    """Adjusts column widths automatically to fit cell contents with padding."""
    for col in range(1, max_cols + 1):
        max_len = 0
        col_letter = get_column_letter(col)
        for row in range(1, ws.max_row + 1):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


def generate_carton_excel(cartons: list[models.Carton], mode: str = "summary") -> bytes:
    """
    Generates an Excel workbook binary from a list of Carton models.
    
    :param cartons: List of Carton instances (with joined product, customer, items)
    :param mode: 'summary' or 'detailed'
    :return: bytes of the .xlsx file
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    if wb.active is not None:
        wb.remove(wb.active)

    # 1. Sheet 1: Tổng hợp Carton
    ws_summary = wb.create_sheet(title="Tong_Hop_Carton")
    ws_summary.views.sheetView[0].showGridLines = True

    summary_headers = [
        "STT",
        "Mã Carton SN",
        "Thời Gian Đóng",
        "Khách Hàng",
        "Sản Phẩm",
        "Chế Độ Đóng Gói",
        "Job Order / PO",
        "Lot#",
        "Trọng Lượng (kg)",
        "Trạng Thái",
        "In Lại",
        "Trạm In"
    ]
    style_header_row(ws_summary, summary_headers, header_fill_color="1E1B4B")

    regular_font = Font(name="Arial", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    success_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    failed_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    for idx, carton in enumerate(cartons, start=1):
        created_str = carton.created_at.strftime("%Y-%m-%d %H:%M:%S") if carton.created_at else ""
        cust_name = carton.product.customer.name if carton.product and carton.product.customer else ""
        prod_name = carton.product.item_name if carton.product else ""
        packing_mode_str = "Đóng gói theo cân" if carton.product and carton.product.packing_mode == "weight_scale" else "Quét mã con"
        order_or_po = carton.job_order or carton.po_number or ""
        lot_str = carton.lot_number or ""
        weight_val = f"{carton.weight:.3f}" if carton.weight is not None else ""
        status_val = carton.status or "UNKNOWN"
        reprint_str = "Có" if getattr(carton, "is_reprint", 0) == 1 else "Không"
        station_str = carton.station_id or ""

        row_values = [
            idx,
            carton.carton_sn,
            created_str,
            cust_name,
            prod_name,
            packing_mode_str,
            order_or_po,
            lot_str,
            weight_val,
            status_val,
            reprint_str,
            station_str
        ]
        ws_summary.append(row_values)
        row_num = idx + 1
        ws_summary.row_dimensions[row_num].height = 20

        # Apply cell formatting
        for col_idx in range(1, len(row_values) + 1):
            cell = ws_summary.cell(row=row_num, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            
            # Alignments
            if col_idx in (1, 3, 6, 8, 10, 11, 12):
                cell.alignment = center_align
            elif col_idx in (9,):
                cell.alignment = right_align
            else:
                cell.alignment = left_align

            # Row zebra striping
            if idx % 2 == 0 and col_idx != 10:
                cell.fill = zebra_fill

            # Highlight status column
            if col_idx == 10:
                if status_val == "SUCCESS":
                    cell.fill = success_fill
                elif status_val == "FAILED":
                    cell.fill = failed_fill

    auto_fit_columns(ws_summary, len(summary_headers))

    # 2. Sheet 2: Chi tiết Serial con (Traceability) khi mode == 'detailed'
    if mode == "detailed":
        ws_detail = wb.create_sheet(title="Chi_Tiet_Serial_Con")
        ws_detail.views.sheetView[0].showGridLines = True

        detail_headers = [
            "STT",
            "Mã Carton SN",
            "Mã Sê-ri Con (Item SN)",
            "Khách Hàng",
            "Sản Phẩm",
            "Job Order",
            "Thời Gian Đóng"
        ]
        style_header_row(ws_detail, detail_headers, header_fill_color="312E81")

        detail_idx = 1
        for carton in cartons:
            items = getattr(carton, "items", []) or []
            cust_name = carton.product.customer.name if carton.product and carton.product.customer else ""
            prod_name = carton.product.item_name if carton.product else ""
            created_str = carton.created_at.strftime("%Y-%m-%d %H:%M:%S") if carton.created_at else ""
            job_order_str = carton.job_order or carton.po_number or ""

            if items:
                for item in items:
                    row_vals = [
                        detail_idx,
                        carton.carton_sn,
                        item.item_sn,
                        cust_name,
                        prod_name,
                        job_order_str,
                        created_str
                    ]
                    ws_detail.append(row_vals)
                    r_num = detail_idx + 1
                    ws_detail.row_dimensions[r_num].height = 20

                    for c_idx in range(1, len(row_vals) + 1):
                        c = ws_detail.cell(row=r_num, column=c_idx)
                        c.font = regular_font
                        c.border = thin_border
                        if c_idx in (1, 6, 7):
                            c.alignment = center_align
                        else:
                            c.alignment = left_align
                        if detail_idx % 2 == 0:
                            c.fill = zebra_fill

                    detail_idx += 1
            else:
                # Dành cho thùng weight_scale không có item_sn.
                row_vals = [
                    detail_idx,
                    carton.carton_sn,
                    "(N/A - Đóng gói theo cân)",
                    cust_name,
                    prod_name,
                    job_order_str,
                    created_str
                ]
                ws_detail.append(row_vals)
                r_num = detail_idx + 1
                ws_detail.row_dimensions[r_num].height = 20
                for c_idx in range(1, len(row_vals) + 1):
                    c = ws_detail.cell(row=r_num, column=c_idx)
                    c.font = regular_font
                    c.border = thin_border
                    if c_idx in (1, 6, 7):
                        c.alignment = center_align
                    else:
                        c.alignment = left_align
                    if detail_idx % 2 == 0:
                        c.fill = zebra_fill
                detail_idx += 1

        auto_fit_columns(ws_detail, len(detail_headers))

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
