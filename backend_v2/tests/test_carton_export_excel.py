import io
import datetime
import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from src.core.database import Base, get_db
from src.core import models
from main import app

@pytest.fixture
def test_setup():
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool
    )
    Base.metadata.create_all(bind=engine)
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    db = TestingSessionLocal()

    def override_get_db():
        try:
            yield db
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    client = TestClient(app)

    yield db, client

    app.dependency_overrides.clear()
    db.close()

def test_export_summary_excel(test_setup):
    """Test exporting summary carton data to Excel."""
    db, client = test_setup

    customer = models.Customer(code="UI", name="Universal Instruments")
    db.add(customer)
    db.commit()

    product = models.Product(
        customer_id=customer.id,
        item_name="Cable Test RJ45",
        packing_mode="item_scan",
        packed_qty=10
    )
    db.add(product)
    db.commit()

    c1 = models.Carton(
        product_id=product.id,
        carton_sn="CN26090100001",
        created_at=datetime.datetime(2026, 9, 1, 10, 30, 0),
        status="SUCCESS",
        job_order="JO-5544",
        is_reprint=0,
        station_id="192.168.1.10"
    )
    c2 = models.Carton(
        product_id=product.id,
        carton_sn="CN26090100002",
        created_at=datetime.datetime(2026, 9, 1, 11, 45, 0),
        status="FAILED",
        job_order="JO-5544",
        is_reprint=1,
        station_id="192.168.1.11"
    )
    db.add_all([c1, c2])
    db.commit()

    response = client.get("/api/v1/cartons/export?mode=summary")
    assert response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in response.headers["content-type"]
    assert "carton_export_summary" in response.headers["content-disposition"]

    # Verify Excel workbook content
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet = wb.active
    assert sheet is not None
    assert sheet.title == "Tong_Hop_Carton"

    # Verify headers
    headers = [cell.value for cell in sheet[1]]
    assert "Mã Carton SN" in headers
    assert "Khách Hàng" in headers
    assert "Sản Phẩm" in headers
    assert "Trạng Thái" in headers
    assert "In Lại" in headers

    # Verify rows (2 data rows)
    sn_col_idx = headers.index("Mã Carton SN") + 1
    sns = [sheet.cell(row=r, column=sn_col_idx).value for r in range(2, sheet.max_row + 1)]
    assert "CN26090100001" in sns
    assert "CN26090100002" in sns


def test_export_detailed_traceability_excel(test_setup):
    """Test exporting detailed traceability carton data to 2-sheet Excel."""
    db, client = test_setup

    customer = models.Customer(code="UI", name="Universal Instruments")
    db.add(customer)
    db.commit()

    product_ui = models.Product(
        customer_id=customer.id,
        item_name="RJ45 Cable UI",
        packing_mode="item_scan",
        packed_qty=2
    )
    product_a11 = models.Product(
        customer_id=customer.id,
        item_name="Scale Product A11",
        packing_mode="weight_scale",
        packed_qty=100
    )
    db.add_all([product_ui, product_a11])
    db.commit()

    c_ui = models.Carton(
        product_id=product_ui.id,
        carton_sn="CN-UI-100",
        created_at=datetime.datetime(2026, 9, 1, 10, 0, 0),
        status="SUCCESS",
        job_order="JO-01"
    )
    db.add(c_ui)
    db.commit()

    item1 = models.CartonItem(carton_id=c_ui.id, item_sn="ITEM-SN-001")
    item2 = models.CartonItem(carton_id=c_ui.id, item_sn="ITEM-SN-002")
    db.add_all([item1, item2])

    c_a11 = models.Carton(
        product_id=product_a11.id,
        carton_sn="CN-A11-200",
        created_at=datetime.datetime(2026, 9, 1, 11, 0, 0),
        status="SUCCESS",
        po_number="PO-99",
        weight=12.345
    )
    db.add(c_a11)
    db.commit()

    response = client.get("/api/v1/cartons/export?mode=detailed")
    assert response.status_code == 200
    assert "carton_export_detailed" in response.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet_names = wb.sheetnames
    assert "Tong_Hop_Carton" in sheet_names
    assert "Chi_Tiet_Serial_Con" in sheet_names

    ws_detail = wb["Chi_Tiet_Serial_Con"]
    headers = [cell.value for cell in ws_detail[1]]
    assert "Mã Carton SN" in headers
    assert "Mã Sê-ri Con (Item SN)" in headers

    carton_col = headers.index("Mã Carton SN") + 1
    item_col = headers.index("Mã Sê-ri Con (Item SN)") + 1

    detail_rows = [
        (ws_detail.cell(row=r, column=carton_col).value, ws_detail.cell(row=r, column=item_col).value)
        for r in range(2, ws_detail.max_row + 1)
    ]

    # Verify item_scan carton items
    assert ("CN-UI-100", "ITEM-SN-001") in detail_rows
    assert ("CN-UI-100", "ITEM-SN-002") in detail_rows

    # Verify weight_scale carton entry
    a11_entries = [row for row in detail_rows if row[0] == "CN-A11-200"]
    assert len(a11_entries) == 1
    assert "N/A" in a11_entries[0][1]


def test_history_rbac_permissions(test_setup):
    """Test RBAC on carton operations: QA can export, but blocked from delete/reprint."""
    from src.features.auth.service import seed_default_users

    db, client = test_setup
    seed_default_users(db)

    # QA token
    resp_qa = client.post("/auth/login", json={"username": "qa", "password": "qa123"})
    assert resp_qa.status_code == 200
    qa_headers = {"Authorization": f"Bearer {resp_qa.json()['access_token']}"}

    # Admin token
    resp_admin = client.post("/auth/login", json={"username": "admin", "password": "admin123"})
    assert resp_admin.status_code == 200
    admin_headers = {"Authorization": f"Bearer {resp_admin.json()['access_token']}"}

    # Setup a sample carton
    cust = models.Customer(code="UI_RBAC", name="Universal")
    db.add(cust)
    db.commit()
    prod = models.Product(customer_id=cust.id, item_name="Prod RBAC", packed_qty=5)
    db.add(prod)
    db.commit()
    carton = models.Carton(product_id=prod.id, carton_sn="SN-RBAC-01", status="SUCCESS")
    db.add(carton)
    db.commit()

    # 1. QA can export
    resp_export = client.get("/api/v1/cartons/export?mode=summary", headers=qa_headers)
    assert resp_export.status_code == 200

    # 2. QA cannot delete (403)
    resp_del_qa = client.delete(f"/api/v1/cartons/{carton.id}", headers=qa_headers)
    assert resp_del_qa.status_code == 403

    # 3. QA cannot reprint (403)
    resp_reprint_qa = client.post(f"/api/v1/print/carton/{carton.id}/reprint", headers=qa_headers)
    assert resp_reprint_qa.status_code == 403

    # 4. Admin can delete
    resp_del_admin = client.delete(f"/api/v1/cartons/{carton.id}", headers=admin_headers)
    assert resp_del_admin.status_code == 200


